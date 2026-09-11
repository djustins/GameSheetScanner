#!/usr/bin/env python3
"""
scripts/import_summer_player_list.py

One-time import from the league's admin spreadsheet ("Team Pitt Penguin
Admin Summer 26.xlsx" — kept out of git, see .gitignore's *.xlsx rule,
since it carries player/parent PII) into Postgres:

  - Every player on the "Summer Player List" tab: name, birth date, and
    parent contact info. No address is imported (the schema doesn't track
    one) and neither is position — this app re-derives position per
    division/team elsewhere rather than storing a single fixed one.
    current_division_id is set to the current (2026 Summer) division for
    all of them, since that's what this sheet is a signup list for.
  - A "Summer 25 Rating" value creates/uses a 2025 Summer Penguin (U10)
    division and records the value as that player's evaluation grade
    there; a "Fall Rating" value does the same for a 2025 Fall Penguin
    (U10) division. These are separate from — and don't change — the
    player's current_division_id.
  - Every coach on the "Coaches" tab, and their team assignment from the
    "Teams" tab, applied to the current (2026 Summer) division's teams.

Safe to re-run: matches players/coaches by exact (case-insensitive) name
and divisions by (year, season, age_group) — updates existing rows and
skips duplicate evaluations/team assignments instead of piling up
duplicates. A name that matches more than one existing player is skipped
rather than guessed at, and reported so it can be resolved by hand.

Usage:
    export DATABASE_URL=postgresql://user:password@host:port/dbname?sslmode=require
    python scripts/import_summer_player_list.py "Team Pitt Penguin Admin Summer 26.xlsx"
"""

import argparse
import os
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import game_sheet_core as core  # noqa: E402

CURRENT_YEAR = 2026
CURRENT_SEASON = "Summer"
AGE_GROUP = "Penguin"

# Team -> coach display name, from the workbook's "Teams" tab.
TEAM_COACH = {
    "Blackhawks": "Matthew Secosky",
    "Mammoth": "Bobby Junker",
    "Kings": "Mike Caruso",
    "Blues": "Tom Motta",
    "Avalanche": "Robert Dobson",
    "Stars": "Bryan Kowalski",
    "Wild": "Nicholas Cecchini",
}

_BOGUS_PHONE_VALUES = {"no answer", "n/a", "na", "none", ""}


def _clean(value) -> str | None:
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def _clean_phone(value) -> str | None:
    cleaned = _clean(value)
    if cleaned is None or cleaned.lower() in _BOGUS_PHONE_VALUES:
        return None
    return cleaned


def import_players(conn, ws, current_division_id: int) -> None:
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}

    summer_div_id = core.add_division(conn, 2025, "Summer", AGE_GROUP)
    fall_div_id = core.add_division(conn, 2025, "Fall", AGE_GROUP)

    by_name: dict[str, list[dict]] = {}
    for p in core.list_players(conn, include_deleted=True):
        by_name.setdefault(p["name"].strip().lower(), []).append(p)

    created = updated = skipped_ambiguous = evals_added = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        name = _clean(row[idx["Player Name"]])
        if not name:
            continue

        dob = row[idx["Date Of Birth"]]
        birth_date = dob.strftime("%Y-%m-%d") if hasattr(dob, "strftime") else _clean(dob)
        contact_first = _clean(row[idx["Parent FirstName"]])
        contact_last = _clean(row[idx["Parent LastName"]])
        contact_phone = _clean_phone(row[idx["Primary Contact Telephone"]])
        contact_email = _clean(row[idx["Primary Contact Email"]])
        summer_rating = _clean(row[idx["Summer 25 Rating"]])
        fall_rating = _clean(row[idx["Fall Rating"]])

        matches = by_name.get(name.lower(), [])
        if len(matches) > 1:
            print(f"  SKIPPED (ambiguous — {len(matches)} existing players named {name!r}): not auto-merging")
            skipped_ambiguous += 1
            continue
        elif len(matches) == 1:
            player_id = matches[0]["id"]
            core.update_player(
                conn, player_id, birth_date=birth_date, current_division_id=current_division_id,
                contact_first_name=contact_first, contact_last_name=contact_last,
                contact_phone=contact_phone, contact_email=contact_email,
            )
            updated += 1
        else:
            player_id = core.add_player(
                conn, name, birth_date=birth_date, current_division_id=current_division_id,
                contact_first_name=contact_first, contact_last_name=contact_last,
                contact_phone=contact_phone, contact_email=contact_email,
            )
            by_name.setdefault(name.lower(), []).append({"id": player_id, "name": name})
            created += 1

        for rating, division_id in ((summer_rating, summer_div_id), (fall_rating, fall_div_id)):
            if not rating:
                continue
            already = conn.execute(
                "SELECT id FROM evaluations WHERE player_id = %s AND division_id = %s",
                (player_id, division_id),
            ).fetchone()
            if already:
                continue
            core.add_evaluation(conn, player_id, division_id, None, rating)
            evals_added += 1

    print(f"Players: {created} created, {updated} updated, {skipped_ambiguous} skipped (ambiguous name)")
    print(f"Evaluations added: {evals_added}")


def import_coaches(conn, coaches_ws, current_division_id: int) -> None:
    headers = [c.value for c in coaches_ws[1]]
    name_idx = headers.index("Coach")

    existing = {c["name"].strip().lower(): c["id"] for c in core.list_coaches(conn, include_deleted=True)}
    coach_ids: dict[str, int] = {}
    added = 0

    for row in coaches_ws.iter_rows(min_row=2, values_only=True):
        name = _clean(row[name_idx])
        if not name:
            continue
        key = name.lower()
        if key in existing:
            coach_ids[name] = existing[key]
        else:
            coach_id = core.add_coach(conn, name)
            existing[key] = coach_id
            coach_ids[name] = coach_id
            added += 1
    print(f"Coaches: {added} created, {len(coach_ids) - added} already existed")

    teams = {t["name"].strip().lower(): t["id"] for t in core.list_teams(conn, current_division_id)}
    assigned = 0
    for team_name, coach_name in TEAM_COACH.items():
        team_id = teams.get(team_name.strip().lower())
        coach_id = existing.get(coach_name.strip().lower())
        if not team_id:
            print(f"  SKIPPED: no team named {team_name!r} in the current division")
            continue
        if not coach_id:
            print(f"  SKIPPED: no coach named {coach_name!r}")
            continue
        core.assign_coach_to_team(conn, team_id, coach_id)
        assigned += 1
    print(f"Coach-team assignments applied (idempotent, includes pre-existing ones): {assigned}")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("workbook", help="Path to the admin .xlsx file")
    args = parser.parse_args()

    if not Path(args.workbook).exists():
        sys.exit(f"No such file: {args.workbook}")

    dsn = os.environ.get("DATABASE_URL")
    if not dsn:
        sys.exit("Error: set the DATABASE_URL environment variable first.")
    conn = core.init_db(dsn)

    current_division_id = core.add_division(conn, CURRENT_YEAR, CURRENT_SEASON, AGE_GROUP)

    wb = openpyxl.load_workbook(args.workbook, data_only=True)
    import_players(conn, wb["Summer Player List"], current_division_id)
    import_coaches(conn, wb["Coaches"], current_division_id)


if __name__ == "__main__":
    main()
